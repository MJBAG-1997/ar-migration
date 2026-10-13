"""
Build hfp_signoff_pack_v6.xlsx and hfp_review_pack_v1.xlsx.
Also writes scripts/hfp_review_reader.py.

Rule 9: aggregates/UUIDs only in console. No PII.
"""
import shutil, json, os, re, textwrap
from collections import defaultdict, Counter
from datetime import date
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

DATA_ROOT   = r"C:\dev\AR Migration"
V5_PATH     = os.path.join(DATA_ROOT, "hfp_signoff_pack_v5.xlsx")
V6_PATH     = os.path.join(DATA_ROOT, "hfp_signoff_pack_v6.xlsx")
REVIEW_PATH = os.path.join(DATA_ROOT, "hfp_review_pack_v1.xlsx")
READER_PATH = os.path.join(DATA_ROOT, "scripts", "hfp_review_reader.py")
SNAP_DIR    = os.path.join(DATA_ROOT, "snapshot_20261012T090906Z")

def p(*a): print(*a, flush=True)

# ── helpers ────────────────────────────────────────────────────────────────────

PLACEHOLDER_UPPER = {
    "TBC","NA","N/A","NOTKNOWN","UNKNOWN","NOTONFILE",
    "NOTSET","NONE","NOT SET","NOTON","NOTAPPLICABLE",
}

def is_placeholder(pn_raw):
    if not pn_raw or str(pn_raw).strip() in ("","None","nan"):
        return True
    pn_str = str(pn_raw).strip()
    if re.fullmatch(r"[\-\s\.]+", pn_str):
        return True
    norm = re.sub(r"[\s\-]","",pn_str).upper()
    if norm in PLACEHOLDER_UPPER:
        return True
    if len(norm) <= 5 and not re.match(r"^\d{4,5}$|^[A-Z]{2}\d{3,}$", norm):
        return True
    return False

def pence_to_pounds(amt):
    try:
        return int(str(amt)) / 100.0
    except (TypeError, ValueError):
        return 0.0

PLAN_TYPE_MAP = {
    "investment":"Investment","pension":"Pension","protection_policy":"Protection",
    "group_protection":"Group protection","annuity":"Annuity","fixed_asset":"Asset",
    "cash_account":"Cash account","mortgage":"Mortgage","loan_credit":"Loan / credit",
}

def plan_type_english(t):
    return PLAN_TYPE_MAP.get(t, t or "")

def plan_polno(pl):
    pn = pl.get("policy_number","")
    return "" if is_placeholder(pn) else (pn or "")

def plan_value_pounds_fn(pl):
    val = pl.get("value")
    if isinstance(val, dict):
        return pence_to_pounds(val.get("amount", 0))
    return 0.0

def get_provider_name(pl):
    prov = pl.get("provider")
    if isinstance(prov, dict):
        return prov.get("name","")
    return ""

# ── load snapshot data ─────────────────────────────────────────────────────────

p("Loading snapshot data...")

with open(os.path.join(SNAP_DIR, "plans_dedup.json"), encoding="utf-8") as f:
    plans_raw = json.load(f)
plans = list(plans_raw.values())
plans_by_uuid = {pl["uuid"]: pl for pl in plans}
p(f"  Plans: {len(plans)}")

with open(os.path.join(SNAP_DIR, "charges_dedup.json"), encoding="utf-8") as f:
    charges_raw = json.load(f)
charges = list(charges_raw.values())
p(f"  Charges: {len(charges)}")

with open(os.path.join(SNAP_DIR, "accounts.json"), encoding="utf-8") as f:
    accounts = json.load(f)
p(f"  Accounts: {len(accounts)}")

with open(os.path.join(SNAP_DIR, "circles.json"), encoding="utf-8") as f:
    circles_data = json.load(f)
p(f"  Circles: {len(circles_data)}")

with open(os.path.join(SNAP_DIR, "reference.json"), encoding="utf-8") as f:
    reference = json.load(f)
employees = reference.get("employees", [])
p(f"  Employees: {len(employees)}")

# ── build circles maps ─────────────────────────────────────────────────────────

p("\nBuilding circles maps...")

circle_name_map = {}         # circle_uuid -> name
circle_to_accounts = defaultdict(set)   # circle_uuid -> set of account_uuids
account_to_circles = defaultdict(list)  # account_uuid -> list of circle_uuids

for acc in accounts:
    acc_uuid = acc["uuid"]
    for c in acc.get("circles", []):
        cuuid = c["uuid"]
        cname = c.get("name","")
        circle_name_map[cuuid] = cname
        circle_to_accounts[cuuid].add(acc_uuid)
        account_to_circles[acc_uuid].append(cuuid)

# Supplement circle_name_map from circles.json (catches circles with no accounts)
for c in circles_data:
    cuuid = c["uuid"]
    if cuuid not in circle_name_map:
        circle_name_map[cuuid] = c.get("name","")

accounts_in_circles = sum(1 for a in accounts if a.get("circles"))
accounts_multi_circles = [
    a["uuid"] for a in accounts if len(a.get("circles", [])) > 1
]
p(f"  Distinct circles (from accounts): {len(circle_to_accounts)}")
p(f"  Accounts in >=1 circle: {accounts_in_circles}")
p(f"  Accounts in >1 circle:  {len(accounts_multi_circles)}")

# ── build reference maps ───────────────────────────────────────────────────────

p("\nBuilding reference maps...")

adviser_name_map = {e["uuid"]: e.get("name","") for e in employees}
acc_by_uuid   = {a["uuid"]: a for a in accounts}
acc_name      = {a["uuid"]: a.get("name","") for a in accounts}

# account_uuid -> adviser_uuid  (from assigned_adviser)
acc_adviser = {}
for a in accounts:
    aa = a.get("assigned_adviser")
    if isinstance(aa, dict) and aa.get("uuid"):
        acc_adviser[a["uuid"]] = aa["uuid"]

# plan_uuid -> primary client uuid (highest ownership %)
plan_primary_client = {}
for pl in plans:
    owners = pl.get("_owners", [])
    if isinstance(owners, list) and owners:
        best = max(owners, key=lambda o: float(o.get("ownership_percentage", 0) or 0))
        plan_primary_client[pl["uuid"]] = best.get("client_uuid","")

# plan_uuid -> adviser_uuid  (charges first, then client assigned_adviser)
plan_to_adviser_uuid = {}
for ch in charges:
    s = ch.get("seller")
    pid = ch.get("_plan_uuid")
    if isinstance(s, dict) and s.get("uuid") and pid:
        plan_to_adviser_uuid[pid] = s["uuid"]
for pl in plans:
    uuid = pl["uuid"]
    if uuid in plan_to_adviser_uuid:
        continue
    client_uuid = plan_primary_client.get(uuid,"")
    if client_uuid and client_uuid in acc_adviser:
        plan_to_adviser_uuid[uuid] = acc_adviser[client_uuid]

def plan_adviser_name(uuid):
    adv_uuid = plan_to_adviser_uuid.get(uuid,"")
    return adviser_name_map.get(adv_uuid, "Unassigned")

def plan_household_name(uuid):
    client_uuid = plan_primary_client.get(uuid,"")
    circles_list = account_to_circles.get(client_uuid,[])
    if circles_list:
        return circle_name_map.get(circles_list[0],"")
    return acc_name.get(client_uuid,"")

def plan_client_name(uuid):
    client_uuid = plan_primary_client.get(uuid,"")
    return acc_name.get(client_uuid,"")

no_adviser = sum(1 for pl in plans if pl.get("status")=="active" and pl["uuid"] not in plan_to_adviser_uuid)
no_client  = sum(1 for pl in plans if pl.get("status")=="active" and not plan_client_name(pl["uuid"]))
p(f"  Active plans with no adviser mapping: {no_adviser}")
p(f"  Active plans with no client name:     {no_client}")

# ── load v5 tabs ───────────────────────────────────────────────────────────────

p("\nLoading v5 tabs for section data...")
df_ca_v5  = pd.read_excel(V5_PATH, sheet_name="Charges (all)", dtype=str)
df_cs_v5  = pd.read_excel(V5_PATH, sheet_name="Charges sign-off", dtype=str)
df_ex_v5  = pd.read_excel(V5_PATH, sheet_name="Exceptions", dtype=str)
df_pop_v5 = pd.read_excel(V5_PATH, sheet_name="Provider-only plans", dtype=str)
p(f"  Charges (all): {len(df_ca_v5)}")
p(f"  Charges sign-off: {len(df_cs_v5)}")
p(f"  Exceptions: {len(df_ex_v5)}")
p(f"  Provider-only: {len(df_pop_v5)}")

# Section sets from v5 Charges sign-off
sec_a_rows_v5 = df_cs_v5[df_cs_v5.get("source_section","").eq("A - Workbook confirmed") if "source_section" in df_cs_v5.columns else df_cs_v5.index == -1]
sec_b_rows_v5 = df_cs_v5[df_cs_v5.get("source_section","").eq("B - Plannr ongoing, no workbook confirm") if "source_section" in df_cs_v5.columns else df_cs_v5.index == -1]
sec_c_rows_v5 = df_cs_v5[df_cs_v5.get("source_section","").eq("C - Under advice, no charge any source") if "source_section" in df_cs_v5.columns else df_cs_v5.index == -1]

if "source_section" in df_cs_v5.columns:
    sec_a_rows_v5 = df_cs_v5[df_cs_v5["source_section"] == "A - Workbook confirmed"]
    sec_b_rows_v5 = df_cs_v5[df_cs_v5["source_section"] == "B - Plannr ongoing, no workbook confirm"]
    sec_c_rows_v5 = df_cs_v5[df_cs_v5["source_section"] == "C - Under advice, no charge any source"]
else:
    sec_a_rows_v5 = sec_b_rows_v5 = sec_c_rows_v5 = pd.DataFrame()

sec_a_plan_uuids_v5 = set(sec_a_rows_v5["plan_uuid"].dropna()) if "plan_uuid" in sec_a_rows_v5.columns else set()
sec_b_plan_uuids_v5 = set(sec_b_rows_v5["plan_uuid"].dropna()) if "plan_uuid" in sec_b_rows_v5.columns else set()
sec_c_plan_uuids_v5 = set(sec_c_rows_v5["plan_uuid"].dropna()) if "plan_uuid" in sec_c_rows_v5.columns else set()

p(f"  Section A plans: {len(sec_a_plan_uuids_v5)}")
p(f"  Section B plans: {len(sec_b_plan_uuids_v5)}")
p(f"  Section C plans: {len(sec_c_plan_uuids_v5)}")

# Section C sub-type lookups
sec_c_cs_map = {}  # plan_uuid -> row dict
if "plan_uuid" in sec_c_rows_v5.columns:
    for _, row in sec_c_rows_v5.iterrows():
        uid = str(row.get("plan_uuid","") or "")
        if uid:
            sec_c_cs_map[uid] = row.to_dict()

def sec_c_why_asking(uuid):
    cs = sec_c_cs_map.get(uuid, {})
    ar_comment  = str(cs.get("AR_comment","") or "")
    val_source  = str(cs.get("valuation_source","") or "")
    if "Rathbones" in val_source and not ar_comment:
        return "This plan is with Rathbones but we have no fee — please give us the rate"
    if "zero" in ar_comment.lower():
        return "The fee on the provider statement is zero — please confirm there is no fee, or give us the rate"
    return "We have no fee for this plan — please confirm there is no fee, or give us the rate"

# Plan charge info from v5 Charges (all)
plan_charge_v5 = {}
if "plan_uuid" in df_ca_v5.columns:
    for _, row in df_ca_v5.iterrows():
        uuid = str(row.get("plan_uuid","") or "")
        if not uuid:
            continue
        try:
            pct = float(row.get("percentage",0) or 0)
        except (ValueError, TypeError):
            pct = 0.0
        try:
            iai = float(row.get("indicative_annual_income",0) or 0)
        except (ValueError, TypeError):
            iai = 0.0
        plan_charge_v5[uuid] = {
            "percentage": pct,          # percentage form: 0.75 = 0.75%/yr
            "pct_decimal": pct / 100.0, # decimal form for Excel 0.00% format
            "iai": iai,
            "adviser_uuid": str(row.get("adviser_uuid","") or ""),
        }

def fee_source(uuid):
    if uuid in sec_a_plan_uuids_v5:
        cs = sec_a_rows_v5[sec_a_rows_v5["plan_uuid"] == uuid] if "plan_uuid" in sec_a_rows_v5.columns else pd.DataFrame()
        if not cs.empty:
            basis = str(cs.iloc[0].get("basis","") or "")
            if "Rathbones" in basis:
                return "Rathbones statement"
        return "Provider statement"
    if uuid in sec_b_plan_uuids_v5:
        return "Your Plannr"
    if uuid in sec_c_plan_uuids_v5:
        return "To confirm"
    return "No fee"

# ── openpyxl style helpers ─────────────────────────────────────────────────────

HEADER_FILL = PatternFill(fill_type="solid", fgColor="F2F2F2")
HEADER_FONT = Font(bold=True)
CURRENCY_FMT = '\u00a3#,##0.00'   # £#,##0.00
PCT_FMT      = '0.00%'

def style_header_row(ws):
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    ws.freeze_panes = "A2"

def apply_autofilter(ws):
    if ws.max_row > 1 and ws.max_column >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"

def set_col_widths_auto(ws, min_w=12, max_w=40):
    for col_cells in ws.columns:
        max_len = min_w
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            if cell.value is not None:
                try:
                    max_len = max(max_len, min(max_w, len(str(cell.value)) + 2))
                except Exception:
                    pass
        ws.column_dimensions[col_letter].width = max_len

def hide_col(ws, col_idx):
    letter = get_column_letter(col_idx)
    ws.column_dimensions[letter].hidden = True
    ws.column_dimensions[letter].width = 0.1

def write_header(ws, headers):
    for i, h in enumerate(headers, start=1):
        ws.cell(row=1, column=i, value=h)
    style_header_row(ws)

def add_dropdown(ws, col_letter, start_row, end_row, options):
    formula = '"' + ",".join(options) + '"'
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv)
    dv.sqref = f"{col_letter}{start_row}:{col_letter}{end_row}"

# ═══════════════════════════════════════════════════════════════════════════════
# DELIVERABLE 1 — hfp_signoff_pack_v6.xlsx
# ═══════════════════════════════════════════════════════════════════════════════

p(f"\nCopying v5 -> v6...")
shutil.copy2(V5_PATH, V6_PATH)

p("Opening v6 with openpyxl...")
wb6 = openpyxl.load_workbook(V6_PATH)

# ── Circles tab ───────────────────────────────────────────────────────────────

p("Adding Circles tab...")
cl_idx = wb6.sheetnames.index("Change log")
ws_circles = wb6.create_sheet("Circles", cl_idx)

circ_headers = ["circle_uuid","circle_name","account_uuid","account_source_uuid","member_count","group"]
write_header(ws_circles, circ_headers)

circle_rows_written = 0
for acc in accounts:
    acc_uuid = acc["uuid"]
    for c in acc.get("circles",[]):
        cuuid = c["uuid"]
        cname = c.get("name","")
        member_cnt = len(circle_to_accounts.get(cuuid, set()))
        row_data = [cuuid, cname, acc_uuid, acc_uuid, member_cnt, "default"]
        ws_circles.append(row_data)
        circle_rows_written += 1

set_col_widths_auto(ws_circles)
apply_autofilter(ws_circles)
p(f"  Circles tab rows: {circle_rows_written}")

# ── Circles exceptions ────────────────────────────────────────────────────────

p("Appending circles exceptions to Exceptions tab...")
ws_ex6 = wb6["Exceptions"]
ex_hdrs = [c.value for c in ws_ex6[1]]
ex_last_row = ws_ex6.max_row

circles_exc_added = 0
for acc_uuid in accounts_multi_circles:
    n = len(account_to_circles[acc_uuid])
    row_data = {
        "category":      "Account in multiple circles",
        "entity_type":   "account",
        "source_uuid":   acc_uuid,
        "detail":        f"Account appears in {n} circles",
        "current_value": str(n),
        "proposed_value":"",
        "AR_decision":   "",
    }
    write_row = [row_data.get(h,"") for h in ex_hdrs]
    ws_ex6.append(write_row)
    circles_exc_added += 1

p(f"  Circles exceptions appended: {circles_exc_added}")

# Update Summary tab version
ws_sum6 = wb6["Summary"]
for row in ws_sum6.iter_rows():
    for cell in row:
        if str(cell.value or "") == "Pack version":
            ws_sum6.cell(row=cell.row, column=cell.column+1, value="v6")
        if str(cell.value or "") == "Build date":
            ws_sum6.cell(row=cell.row, column=cell.column+1, value=str(date.today()))

p(f"Saving v6 to {V6_PATH}...")
wb6.save(V6_PATH)
p("  v6 saved.")

# ═══════════════════════════════════════════════════════════════════════════════
# DELIVERABLE 2 — hfp_review_pack_v1.xlsx
# ═══════════════════════════════════════════════════════════════════════════════

p("\nBuilding hfp_review_pack_v1.xlsx from scratch...")
wbr = openpyxl.Workbook()
wbr.remove(wbr.active)

# ── Tab 1: Read me ─────────────────────────────────────────────────────────────

p("  Tab: Read me")
ws_rm = wbr.create_sheet("Read me")
lines = [
    "1. This document is an adviser review pack for Hoban Financial Planning Ltd. It is for the adviser's eyes only and contains a summary of clients, plans and fees as at the migration date.",
    "2. Three tabs need your action before we can finalise the migration: 'Fees to confirm' (check or update ongoing fees), 'Plans to check' (resolve flagged issues with policy numbers and Rathbones plans), and 'Plans we found at providers' (confirm which clients these belong to).",
    "3. To fill in answers: use the dropdown in the 'Your answer' column on each tab. Where you see a 'Comment' or 'New fee %' column, add any extra detail there.",
    "4. Please do not add or remove rows, or rename column headers — the reader script uses them to import your responses.",
    "5. Return by: ___________",
]
for i, line in enumerate(lines, start=1):
    ws_rm.cell(row=i, column=1, value=line)
ws_rm.column_dimensions["A"].width = 110

# ── Tab 2: Summary ─────────────────────────────────────────────────────────────

p("  Tab: Summary")
ws_sum = wbr.create_sheet("Summary")

from collections import defaultdict as _dd
adv_clients  = _dd(set)
adv_circles  = _dd(set)
adv_uncircled= _dd(set)
adv_plans    = _dd(int)
adv_val_ua   = _dd(float)

active_plans = [pl for pl in plans if pl.get("status")=="active"]
for pl in active_plans:
    uuid     = pl["uuid"]
    adv_uuid = plan_to_adviser_uuid.get(uuid,"unassigned")
    c_uuid   = plan_primary_client.get(uuid,"")
    adv_plans[adv_uuid] += 1
    if c_uuid:
        adv_clients[adv_uuid].add(c_uuid)
        circ_list = account_to_circles.get(c_uuid,[])
        if circ_list:
            adv_circles[adv_uuid].update(circ_list)
        else:
            adv_uncircled[adv_uuid].add(c_uuid)
    if pl.get("under_advice") in ("Yes", True) or pl.get("under_advice") == "true":
        adv_val_ua[adv_uuid] += plan_value_pounds_fn(pl)

# IAI from v5 Charges (all)
adv_iai = _dd(float)
if "adviser_uuid" in df_ca_v5.columns and "indicative_annual_income" in df_ca_v5.columns:
    for _, row in df_ca_v5.iterrows():
        adv_uuid = str(row.get("adviser_uuid","") or "")
        try:
            iai = float(row.get("indicative_annual_income",0) or 0)
        except (ValueError, TypeError):
            iai = 0.0
        if adv_uuid:
            adv_iai[adv_uuid] += iai

sum_headers = ["Adviser","Clients","Households","Active plans",
               "Value under advice \u00a3","Ongoing fee \u00a3/yr","Avg fee %"]
write_header(ws_sum, sum_headers)

all_adv_uuids = set(adv_plans.keys())
sum_rows = []
for adv_uuid in sorted(all_adv_uuids, key=lambda u: adviser_name_map.get(u,"Unassigned")):
    name       = adviser_name_map.get(adv_uuid,"Unassigned")
    clients    = len(adv_clients[adv_uuid])
    households = len(adv_circles[adv_uuid]) + len(adv_uncircled[adv_uuid])
    n_plans    = adv_plans[adv_uuid]
    val_ua     = adv_val_ua[adv_uuid]
    iai        = adv_iai[adv_uuid]
    avg_fee    = iai / val_ua if val_ua > 0 else 0.0
    sum_rows.append([name, clients, households, n_plans, val_ua, iai, avg_fee])

total_clients = len(set().union(*adv_clients.values()) if adv_clients else set())
total_hh = sum(len(adv_circles[u]) + len(adv_uncircled[u]) for u in all_adv_uuids)
total_plans = sum(adv_plans.values())
total_val = sum(adv_val_ua.values())
total_iai = sum(adv_iai.values())
sum_rows.append(["TOTAL", total_clients, total_hh, total_plans, total_val, total_iai,
                 total_iai/total_val if total_val > 0 else 0.0])

for r_idx, row_data in enumerate(sum_rows, start=2):
    for c_idx, val in enumerate(row_data, start=1):
        ws_sum.cell(row=r_idx, column=c_idx, value=val)

# Apply formats
val_ua_col  = get_column_letter(5)
iai_col     = get_column_letter(6)
avg_fee_col = get_column_letter(7)
for r in range(2, len(sum_rows)+2):
    ws_sum[f"{val_ua_col}{r}"].number_format  = CURRENCY_FMT
    ws_sum[f"{iai_col}{r}"].number_format     = CURRENCY_FMT
    ws_sum[f"{avg_fee_col}{r}"].number_format = PCT_FMT

set_col_widths_auto(ws_sum)

# ── Tab 3: Client book ─────────────────────────────────────────────────────────

p("  Tab: Client book")
ws_cb = wbr.create_sheet("Client book")

cb_headers = ["Adviser","Household","Client","Plan type","Provider",
              "Policy no","Value \u00a3","Under advice","Ongoing fee %","Fee source","_ref"]
write_header(ws_cb, cb_headers)

cb_rows = []
for pl in active_plans:
    uuid = pl["uuid"]
    pct_info = plan_charge_v5.get(uuid)
    ongoing_pct = pct_info["pct_decimal"] if pct_info else None
    cb_rows.append({
        "Adviser":       plan_adviser_name(uuid),
        "Household":     plan_household_name(uuid),
        "Client":        plan_client_name(uuid),
        "Plan type":     plan_type_english(pl.get("type","")),
        "Provider":      get_provider_name(pl),
        "Policy no":     plan_polno(pl),
        "Value \u00a3":  plan_value_pounds_fn(pl),
        "Under advice":  pl.get("under_advice",""),
        "Ongoing fee %": ongoing_pct,
        "Fee source":    fee_source(uuid),
        "_ref":          uuid,
    })

cb_rows.sort(key=lambda r: (r["Adviser"], r["Household"], r["Client"], r["Plan type"]))

val_col_cb  = get_column_letter(cb_headers.index("Value \u00a3") + 1)
pct_col_cb  = get_column_letter(cb_headers.index("Ongoing fee %") + 1)
ref_col_cb  = len(cb_headers)

for r_idx, row_data in enumerate(cb_rows, start=2):
    for c_idx, h in enumerate(cb_headers, start=1):
        ws_cb.cell(row=r_idx, column=c_idx, value=row_data.get(h))
    ws_cb[f"{val_col_cb}{r_idx}"].number_format = CURRENCY_FMT
    val_pct = row_data["Ongoing fee %"]
    if val_pct is not None:
        ws_cb[f"{pct_col_cb}{r_idx}"].number_format = PCT_FMT

apply_autofilter(ws_cb)
set_col_widths_auto(ws_cb)
hide_col(ws_cb, ref_col_cb)
p(f"    Client book rows: {len(cb_rows)}")

# ── Tab 4: Fees to confirm ─────────────────────────────────────────────────────

p("  Tab: Fees to confirm")
ws_ftc = wbr.create_sheet("Fees to confirm")

ftc_headers = ["Client","Plan type","Provider","Policy no","Value \u00a3",
               "Fee we have %","Why we are asking","Your answer","New fee %","_ref"]
write_header(ws_ftc, ftc_headers)

ftc_rows = []
for uuid in sorted(sec_b_plan_uuids_v5 | sec_c_plan_uuids_v5):
    pl = plans_by_uuid.get(uuid)
    if not pl:
        continue
    pct_info = plan_charge_v5.get(uuid)
    fee_pct  = pct_info["pct_decimal"] if pct_info else None
    if uuid in sec_b_plan_uuids_v5:
        why = "We have a fee from your Plannr — please confirm it is still correct"
    else:
        why = sec_c_why_asking(uuid)
    ftc_rows.append({
        "Client":          plan_client_name(uuid),
        "Plan type":       plan_type_english(pl.get("type","")),
        "Provider":        get_provider_name(pl),
        "Policy no":       plan_polno(pl),
        "Value \u00a3":    plan_value_pounds_fn(pl),
        "Fee we have %":   fee_pct,
        "Why we are asking": why,
        "Your answer":     None,
        "New fee %":       None,
        "_ref":            uuid,
    })

val_col_ftc = get_column_letter(ftc_headers.index("Value \u00a3") + 1)
fee_col_ftc = get_column_letter(ftc_headers.index("Fee we have %") + 1)
ans_col_ftc = get_column_letter(ftc_headers.index("Your answer") + 1)
ref_col_ftc = len(ftc_headers)

for r_idx, row_data in enumerate(ftc_rows, start=2):
    for c_idx, h in enumerate(ftc_headers, start=1):
        ws_ftc.cell(row=r_idx, column=c_idx, value=row_data.get(h))
    ws_ftc[f"{val_col_ftc}{r_idx}"].number_format = CURRENCY_FMT
    if row_data["Fee we have %"] is not None:
        ws_ftc[f"{fee_col_ftc}{r_idx}"].number_format = PCT_FMT

if ftc_rows:
    add_dropdown(ws_ftc, ans_col_ftc, 2, len(ftc_rows)+1,
                 ["Correct","Change to","No fee"])

apply_autofilter(ws_ftc)
set_col_widths_auto(ws_ftc)
hide_col(ws_ftc, ref_col_ftc)
p(f"    Fees to confirm rows: {len(ftc_rows)}")

# ── Tab 5: Plans to check ──────────────────────────────────────────────────────

p("  Tab: Plans to check")
ws_ptc = wbr.create_sheet("Plans to check")

ptc_headers = ["Client","Plan type","Provider","Policy no","Issue",
               "Your answer","Comment","_ref"]
write_header(ws_ptc, ptc_headers)

PLANS_TO_CHECK_CATS = {
    "No policy number",
    "Duplicate policy number",
    "Rathbones plan not on Rathbones list",
    "Active agency plan not in workbook",
    "Rathbones file: matched to non-Rathbones plan",
}

ISSUE_MAP = {
    "No policy number":
        "We have no policy number for this plan — please add one",
    "Duplicate policy number":
        "This policy number appears on more than one client — please check",
    "Rathbones plan not on Rathbones list":
        "This Rathbones plan was not on the Rathbones list we received — closed or transferred?",
    "Active agency plan not in workbook":
        "This plan is active but was not in the provider statement — please check",
    "Rathbones file: matched to non-Rathbones plan":
        "This plan was matched in the Rathbones file but has a different provider in Plannr — please investigate",
}

ptc_rows = []
if "category" in df_ex_v5.columns and "entity_type" in df_ex_v5.columns:
    ex_plan_rows = df_ex_v5[
        (df_ex_v5["entity_type"] == "plan") &
        (df_ex_v5["category"].isin(PLANS_TO_CHECK_CATS))
    ]
    for _, erow in ex_plan_rows.iterrows():
        uuid = str(erow.get("source_uuid","") or "")
        pl   = plans_by_uuid.get(uuid, {})
        cat  = str(erow.get("category","") or "")
        ptc_rows.append({
            "Client":      plan_client_name(uuid) if pl else "",
            "Plan type":   plan_type_english(pl.get("type","")) if pl else "",
            "Provider":    get_provider_name(pl) if pl else "",
            "Policy no":   plan_polno(pl) if pl else "",
            "Issue":       ISSUE_MAP.get(cat, cat),
            "Your answer": None,
            "Comment":     None,
            "_ref":        uuid,
        })

ans_col_ptc = get_column_letter(ptc_headers.index("Your answer") + 1)
ref_col_ptc = len(ptc_headers)

for r_idx, row_data in enumerate(ptc_rows, start=2):
    for c_idx, h in enumerate(ptc_headers, start=1):
        ws_ptc.cell(row=r_idx, column=c_idx, value=row_data.get(h))

if ptc_rows:
    add_dropdown(ws_ptc, ans_col_ptc, 2, len(ptc_rows)+1,
                 ["Add policy number","No policy number","Remove plan",
                  "Same client - merge","Error - remove one",
                  "Closed","Transferred","Still active","Correct"])

apply_autofilter(ws_ptc)
set_col_widths_auto(ws_ptc)
hide_col(ws_ptc, ref_col_ptc)
p(f"    Plans to check rows: {len(ptc_rows)}")

# ── Tab 6: Plans we found at providers ────────────────────────────────────────

p("  Tab: Plans we found at providers")
ws_prov = wbr.create_sheet("Plans we found at providers")

prov_headers = ["Provider","Policy no","Name on provider record","Value \u00a3",
                "Fee %","Our best match","Your answer","Which client","_ref"]
write_header(ws_prov, prov_headers)

prov_rows_written = 0
val_col_prov = get_column_letter(prov_headers.index("Value \u00a3") + 1)
fee_col_prov = get_column_letter(prov_headers.index("Fee %") + 1)
ans_col_prov = get_column_letter(prov_headers.index("Your answer") + 1)
ref_col_prov = len(prov_headers)

pop_rows_list = []
for _, row in df_pop_v5.iterrows():
    wv = row.get("workbook_value","")
    wr = row.get("workbook_rate","")
    try:
        wv_f = float(wv) if str(wv).strip() not in ("","nan","None") else None
    except (ValueError, TypeError):
        wv_f = None
    try:
        wr_f = float(wr) if str(wr).strip() not in ("","nan","None") else None
    except (ValueError, TypeError):
        wr_f = None
    polno = str(row.get("policy_number","") or "")
    pop_rows_list.append({
        "Provider":               str(row.get("provider_name","") or ""),
        "Policy no":              polno,
        "Name on provider record":str(row.get("wrapper_name","") or ""),
        "Value \u00a3":           wv_f,
        "Fee %":                  wr_f,
        "Our best match":         "",
        "Your answer":            None,
        "Which client":           None,
        "_ref":                   polno,
    })

for r_idx, row_data in enumerate(pop_rows_list, start=2):
    for c_idx, h in enumerate(prov_headers, start=1):
        ws_prov.cell(row=r_idx, column=c_idx, value=row_data.get(h))
    if row_data["Value \u00a3"] is not None:
        ws_prov[f"{val_col_prov}{r_idx}"].number_format = CURRENCY_FMT
    if row_data["Fee %"] is not None:
        ws_prov[f"{fee_col_prov}{r_idx}"].number_format = PCT_FMT
    prov_rows_written += 1

if pop_rows_list:
    add_dropdown(ws_prov, ans_col_prov, 2, len(pop_rows_list)+1,
                 ["Add to this client","Add to another client","Not ours","Closed"])

apply_autofilter(ws_prov)
set_col_widths_auto(ws_prov)
hide_col(ws_prov, ref_col_prov)
p(f"    Plans we found at providers rows: {prov_rows_written}")

# ── save review pack ───────────────────────────────────────────────────────────

p(f"\nSaving review pack to {REVIEW_PATH}...")
wbr.save(REVIEW_PATH)
p("  Review pack saved.")

# ═══════════════════════════════════════════════════════════════════════════════
# DELIVERABLE 3 — hfp_review_reader.py
# ═══════════════════════════════════════════════════════════════════════════════

p(f"\nWriting reader script to {READER_PATH}...")

READER_SRC = textwrap.dedent(r'''
    #!/usr/bin/env python3
    """
    Read a returned hfp_review_pack_v1.xlsx and report answered rows.
    Usage: python hfp_review_reader.py <path_to_review_pack.xlsx>
    Output: Tab | _ref | answer | detail (TSV to stdout)
    """
    import sys
    import openpyxl

    ACTION_TABS = {
        "Fees to confirm":             {"answer": "Your answer", "extra": "New fee %"},
        "Plans to check":              {"answer": "Your answer", "extra": "Comment"},
        "Plans we found at providers": {"answer": "Your answer", "extra": "Which client"},
    }

    def main():
        if len(sys.argv) < 2:
            print("Usage: hfp_review_reader.py <path_to_review_pack.xlsx>", file=sys.stderr)
            sys.exit(1)
        path = sys.argv[1]
        wb = openpyxl.load_workbook(path, data_only=True)
        for tab, cfg in ACTION_TABS.items():
            if tab not in wb.sheetnames:
                print(f"WARNING: tab '{tab}' not found", file=sys.stderr)
                continue
            ws  = wb[tab]
            hdrs = [c.value for c in ws[1]]
            def ci(name):
                return next((i for i, h in enumerate(hdrs) if h == name), None)
            ans_i = ci(cfg["answer"])
            ref_i = ci("_ref")
            ext_i = ci(cfg["extra"])
            if ans_i is None:
                print(f"WARNING: column '{cfg['answer']}' not found in '{tab}'", file=sys.stderr)
                continue
            for row in ws.iter_rows(min_row=2, values_only=True):
                ans = row[ans_i] if ans_i < len(row) else None
                if not ans or str(ans).strip() == "":
                    continue
                ref = row[ref_i] if ref_i is not None and ref_i < len(row) else ""
                ext = row[ext_i] if ext_i is not None and ext_i < len(row) else ""
                print(f"{tab}\t{ref}\t{ans}\t{ext}")

    if __name__ == "__main__":
        main()
''').lstrip()

with open(READER_PATH, "w", encoding="utf-8") as f:
    f.write(READER_SRC)
p("  Reader script written.")

# ═══════════════════════════════════════════════════════════════════════════════
# CONSOLE REPORT
# ═══════════════════════════════════════════════════════════════════════════════

v6_size     = os.path.getsize(V6_PATH)     // 1024
review_size = os.path.getsize(REVIEW_PATH) // 1024

p()
p("=" * 60)
p("FINAL REPORT (Rule 9: aggregates only)")
p("=" * 60)
p()
p("=== Circles (v6) ===")
p(f"  Total circles in circle_name_map:     {len(circle_name_map)}")
p(f"  Accounts in >=1 circle:               {accounts_in_circles}")
p(f"  Accounts in >1 circle (exceptions):   {len(accounts_multi_circles)}")
p(f"  Circles tab rows written:             {circle_rows_written}")
p(f"  Circles exceptions appended to v6:    {circles_exc_added}")
p()
p("=== Review pack tabs ===")
p(f"  Summary rows (advisers):              {len(sum_rows)-1}  (+ TOTAL row)")
p(f"  Client book rows:                     {len(cb_rows)}")
p(f"  Fees to confirm rows:                 {len(ftc_rows)}")
p(f"    Section B:                          {len(sec_b_plan_uuids_v5)}")
p(f"    Section C:                          {len(sec_c_plan_uuids_v5)}")
p(f"  Plans to check rows:                  {len(ptc_rows)}")
p(f"  Plans we found at providers rows:     {prov_rows_written}")
p()
p("=== Warnings ===")
p(f"  Active plans with no adviser mapping: {no_adviser}")
p(f"  Active plans with no client name:     {no_client}")
p()
p("=== Output files ===")
p(f"  v6:          {V6_PATH}  ({v6_size} KB)")
p(f"  Review pack: {REVIEW_PATH}  ({review_size} KB)")
p(f"  Reader:      {READER_PATH}")
