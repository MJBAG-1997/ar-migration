"""
Quick inspection: v2 tab names + Rathbones profile.
Uses openpyxl read_only with targeted cell reads.
"""
import sys, os, re, json
from collections import Counter, defaultdict
import openpyxl

DATA_ROOT = r"C:\dev\AR Migration"
V2_PATH   = os.path.join(DATA_ROOT, "hfp_signoff_pack_v2.xlsx")
RATH_PATH = os.path.join(DATA_ROOT, "Jim list  - Rathbones.xlsx")
SNAP_DIR  = os.path.join(DATA_ROOT, "snapshot_20261012T090906Z")

# Flush immediately
def p(*a):
    print(*a, flush=True)

# ─── v2: tab names only first ─────────────────────────────────────────────────
p("=== V2 TABS ===")
wb2 = openpyxl.load_workbook(V2_PATH, read_only=True, data_only=True)
tabs = wb2.sheetnames
p(tabs)
p()

for tab in tabs:
    ws = wb2[tab]
    # Just read first row headers
    row1 = [ws.cell(1, c).value for c in range(1, min(ws.max_column or 30, 60) + 1)]
    p(f"  {tab} ({ws.max_row}r x {ws.max_column}c):")
    p(f"    row1={[h for h in row1 if h]}")
    # Check first 5 rows for section markers
    markers = []
    for r in range(1, min(6, (ws.max_row or 5) + 1)):
        v = ws.cell(r, 1).value
        if v:
            markers.append(str(v)[:60])
    p(f"    first col cells: {markers}")
    p()

wb2.close()
