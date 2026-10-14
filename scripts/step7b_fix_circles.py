"""
step7b_fix_circles.py - Fix Circles INFO: Members and recheck Clients
'In more than one household' label (account_uuids uses | separator).
Rule 9: names in Excel only, counts in terminal.
"""
import os, csv
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

DATA_ROOT   = r"C:\dev\AR Migration"
EXPORTS_DIR = os.path.join(DATA_ROOT, "exports_20261014")
MANIFEST    = os.path.join(EXPORTS_DIR, "_manifest.json")

CIRCLES_PATH = os.path.join(DATA_ROOT, "Plannr_Circles.xlsx")
CLIENTS_PATH = os.path.join(DATA_ROOT, "Plannr_Clients.xlsx")

import json
with open(MANIFEST, encoding="utf-8") as f:
    man = json.load(f).get("results", {})
def ep(l): return man.get(l, {}).get("path", "")

with open(ep("client_export"), encoding="utf-8-sig", newline="") as f:
    client_rows = list(csv.DictReader(f))
client_by_uuid = {r["uuid"]: r for r in client_rows}

def s(v): return str(v).strip() if v is not None and str(v).strip() not in ("","None","nan") else ""
def blank(v): return not s(v)

def client_name(cuuid):
    """Rule 9: returns name for Excel only."""
    ce = client_by_uuid.get(s(cuuid), {})
    ent = s(ce.get("entity_name",""))
    if ent: return ent
    return " ".join(x for x in [s(ce.get("first_name","")), s(ce.get("last_name",""))] if x)

HDR_FONT  = Font(bold=True)
HDR_FILL  = PatternFill(fill_type="solid", fgColor="F2F2F2")
INFO_FONT = Font(bold=True, color="1F497D")
INFO_FILL = PatternFill(fill_type="solid", fgColor="DCE6F1")
ORIG_FONT = Font(bold=True, color="7F6000")
ORIG_FILL = PatternFill(fill_type="solid", fgColor="FFF2CC")

def style_hdr(cell):
    h = str(cell.value or "")
    if h.startswith("INFO:"):
        cell.font = INFO_FONT; cell.fill = INFO_FILL
    elif h == "Review" or h.startswith("ORIGINAL_"):
        cell.font = ORIG_FONT; cell.fill = ORIG_FILL
    else:
        cell.font = HDR_FONT; cell.fill = HDR_FILL

# ── Fix Circles ────────────────────────────────────────────────────────────────

wb_c = openpyxl.load_workbook(CIRCLES_PATH, read_only=True, data_only=True)
ws_c = wb_c.worksheets[0]
circ_rows = list(ws_c.iter_rows(values_only=True))
wb_c.close()

hdrs_c = list(circ_rows[0])
data_c = [{hdrs_c[i]: r[i] for i in range(len(hdrs_c))}
          for r in circ_rows[1:] if any(v is not None for v in r)]

acuuids_col = next((h for h in hdrs_c if h == "account_uuids"), None)
circle_uuid_col = next((h for h in hdrs_c if h in ("uuid","circle_uuid")), None)

# Build client→circles map for "in more than one household" check
client_circles = defaultdict(list)   # account_uuid → [circle_uuids]
for row in data_c:
    cuuid_str = s(row.get(acuuids_col,"") if acuuids_col else "")
    circle_id = s(row.get(circle_uuid_col,"") if circle_uuid_col else "")
    for cuuid in cuuid_str.split("|"):
        cuuid = cuuid.strip()
        if cuuid and circle_id:
            client_circles[cuuid].append(circle_id)

multi_household = {cuuid for cuuid, circles in client_circles.items() if len(set(circles)) > 1}
print(f"Clients in >1 household (circle): {len(multi_household)}")

# Fix INFO: Members - populate from account_uuids with | separator
m_idx = hdrs_c.index("INFO: Members") if "INFO: Members" in hdrs_c else None
if m_idx is not None:
    for row in data_c:
        cuuid_str = s(row.get(acuuids_col,"") if acuuids_col else "")
        uuids = [u.strip() for u in cuuid_str.split("|") if u.strip()]
        names = [client_name(u) for u in uuids if s(u)]
        row["INFO: Members"] = "; ".join(n for n in names if n)

# Sort circles by name
name_col = next((h for h in hdrs_c if h.lower() == "name"), None)
if name_col:
    data_c.sort(key=lambda r: s(r.get(name_col,"")).lower())

# Write Circles
wb_c2 = openpyxl.load_workbook(CIRCLES_PATH)
idx = wb_c2.sheetnames.index("Circles") if "Circles" in wb_c2.sheetnames else 0
del wb_c2["Circles"]
ws_new = wb_c2.create_sheet("Circles", idx)
for i, h in enumerate(hdrs_c, 1):
    c = ws_new.cell(1, i, h)
    style_hdr(c)
for ri, row in enumerate(data_c, 2):
    for ci, h in enumerate(hdrs_c, 1):
        ws_new.cell(ri, ci, row.get(h))
# Freeze through INFO: Members (col 1)
ws_new.freeze_panes = "B2"
if data_c:
    ws_new.auto_filter.ref = f"A1:{get_column_letter(len(hdrs_c))}{len(data_c)+1}"
wb_c2.save(CIRCLES_PATH)

members_filled = sum(1 for r in data_c if not blank(r.get("INFO: Members","")))
print(f"Circles INFO: Members filled: {members_filled}/{len(data_c)}")

# ── Fix Clients "In more than one household" ────────────────────────────────────

wb_cl = openpyxl.load_workbook(CLIENTS_PATH, read_only=True, data_only=True)
ws_cl = wb_cl["Clients"]
cl_rows = list(ws_cl.iter_rows(values_only=True))
wb_cl.close()

hdrs_cl = list(cl_rows[0])
data_cl = [{hdrs_cl[i]: r[i] for i in range(len(hdrs_cl))}
           for r in cl_rows[1:] if any(v is not None for v in r)]

rev_idx = hdrs_cl.index("Review") if "Review" in hdrs_cl else None
uuid_idx = next((hdrs_cl.index(h) for h in hdrs_cl
                 if h in ("uuid","account_uuid","external_reference_value") and h == "uuid"), None)

# If uuid column is missing at front, find it
if uuid_idx is None:
    uuid_idx = next((i for i,h in enumerate(hdrs_cl) if h == "uuid"), None)

changed = 0
for row in data_cl:
    cuuid = s(row.get("uuid",""))
    if not cuuid:
        # try external_reference_value
        cuuid = s(row.get("external_reference_value",""))

    old_review = s(row.get("Review",""))
    labels = [l.strip() for l in old_review.split(";") if l.strip()]

    was_multi = "In more than one household" in labels
    is_multi  = cuuid in multi_household

    if was_multi and not is_multi:
        labels = [l for l in labels if l != "In more than one household"]
        row["Review"] = "; ".join(labels)
        changed += 1
    elif is_multi and not was_multi:
        labels.append("In more than one household")
        row["Review"] = "; ".join(labels)
        changed += 1

print(f"Clients 'In more than one household' now: {sum(1 for r in data_cl if 'In more than one household' in s(r.get('Review','')))}")
print(f"Client Review rows changed: {changed}")

# Write Clients (preserve sort order from step7 agent)
wb_cl2 = openpyxl.load_workbook(CLIENTS_PATH)
idx_cl = wb_cl2.sheetnames.index("Clients") if "Clients" in wb_cl2.sheetnames else 0
del wb_cl2["Clients"]
ws_cl_new = wb_cl2.create_sheet("Clients", idx_cl)

# Determine freeze col (last INFO or moved col)
moved_cols = {"title","first_name","last_name","entity_name"}
info_or_moved = []
for i, h in enumerate(hdrs_cl, 1):
    h_s = str(h or "")
    if h_s.startswith("INFO:") or h_s in moved_cols:
        info_or_moved.append(i)
freeze_after = max(info_or_moved) if info_or_moved else 0

for i, h in enumerate(hdrs_cl, 1):
    c = ws_cl_new.cell(1, i, h)
    style_hdr(c)
for ri, row in enumerate(data_cl, 2):
    for ci, h in enumerate(hdrs_cl, 1):
        ws_cl_new.cell(ri, ci, row.get(h))
ws_cl_new.freeze_panes = f"{get_column_letter(freeze_after + 1)}2"
if data_cl:
    ws_cl_new.auto_filter.ref = f"A1:{get_column_letter(len(hdrs_cl))}{len(data_cl)+1}"
wb_cl2.save(CLIENTS_PATH)
print(f"Saved Plannr_Clients.xlsx ({len(data_cl)} rows)")

# Final Review counts for Clients
from collections import Counter
review_final = Counter()
for row in data_cl:
    for lbl in s(row.get("Review","")).split(";"):
        lbl = lbl.strip()
        if lbl:
            review_final[lbl] += 1
print("\nClients Review labels (final):")
for lbl, cnt in sorted(review_final.items(), key=lambda x: -x[1]):
    print(f"  {lbl}: {cnt}")
print("STOP")
