#!/usr/bin/env python3
"""
Read a returned hfp_review_pack_v1.xlsx and report answered rows.
Usage: python hfp_review_reader.py <path_to_review_pack.xlsx>
Output: Tab | _ref | answer | detail (TSV to stdout)
"""
import sys
import openpyxl

ACTION_TABS = {
    "Fees to confirm":             {"answer": "Your answer", "extra": "New fee %"},
    "Plans to check":              {"answer": "Your answer", "extra": "Comment"},
    "Plans we found at providers": {"answer": "Your answer", "extra": "Which client"},
}

def main():
    if len(sys.argv) < 2:
        print("Usage: hfp_review_reader.py <path_to_review_pack.xlsx>", file=sys.stderr)
        sys.exit(1)
    path = sys.argv[1]
    wb = openpyxl.load_workbook(path, data_only=True)
    for tab, cfg in ACTION_TABS.items():
        if tab not in wb.sheetnames:
            print(f"WARNING: tab '{tab}' not found", file=sys.stderr)
            continue
        ws  = wb[tab]
        hdrs = [c.value for c in ws[1]]
        def ci(name):
            return next((i for i, h in enumerate(hdrs) if h == name), None)
        ans_i = ci(cfg["answer"])
        ref_i = ci("_ref")
        ext_i = ci(cfg["extra"])
        if ans_i is None:
            print(f"WARNING: column '{cfg['answer']}' not found in '{tab}'", file=sys.stderr)
            continue
        for row in ws.iter_rows(min_row=2, values_only=True):
            ans = row[ans_i] if ans_i < len(row) else None
            if not ans or str(ans).strip() == "":
                continue
            ref = row[ref_i] if ref_i is not None and ref_i < len(row) else ""
            ext = row[ext_i] if ext_i is not None and ext_i < len(row) else ""
            print(f"{tab}\t{ref}\t{ans}\t{ext}")

if __name__ == "__main__":
    main()
